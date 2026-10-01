"""Sanrathi export watermark: stamps a small logo + 'Made with Sanrathi' on
every slide of a PPTX, or every page of a PDF. Run inside the Sanrathi gateway
after downloading the file from the Slide Maker."""
import io
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

LOGO = Path(__file__).with_name("sanrathi-mark.png")   # square PNG, transparent bg
LABEL = "Made with Sanrathi"
INK = RGBColor(0x0C, 0x0C, 0x0B)


def stamp_pptx(src: str, dst: str) -> None:
    prs = Presentation(src)
    w, h = prs.slide_width, prs.slide_height
    size = int(h * 0.045)                  # logo = 4.5% of slide height
    margin = int(h * 0.025)
    for slide in prs.slides:
        x = w - margin - size
        y = h - margin - size
        slide.shapes.add_picture(str(LOGO), Emu(x), Emu(y), Emu(size), Emu(size))
        box = slide.shapes.add_textbox(Emu(x - int(size * 5.2)), Emu(y),
                                       Emu(int(size * 5)), Emu(size))
        tf = box.text_frame
        tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = LABEL
        p.font.size = Pt(9)
        p.font.color.rgb = INK
        p.alignment = PP_ALIGN.RIGHT
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    prs.save(dst)


def stamp_pdf(src: str, dst: str) -> None:
    reader, writer = PdfReader(src), PdfWriter()
    logo = ImageReader(str(LOGO))
    for page in reader.pages:
        pw, ph = float(page.mediabox.width), float(page.mediabox.height)
        size, margin = ph * 0.045, ph * 0.025
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=(pw, ph))
        c.drawImage(logo, pw - margin - size, margin, size, size, mask="auto")
        c.setFont("Helvetica", max(6, size * 0.38))
        c.setFillColorRGB(0.05, 0.05, 0.04)
        c.drawRightString(pw - margin - size * 1.2, margin + size * 0.35, LABEL)
        c.save()
        buf.seek(0)
        page.merge_page(PdfReader(buf).pages[0])
        writer.add_page(page)
    with open(dst, "wb") as f:
        writer.write(f)


def stamp(src: str, dst: str) -> None:
    if src.lower().endswith(".pptx"):
        stamp_pptx(src, dst)
    elif src.lower().endswith(".pdf"):
        stamp_pdf(src, dst)
    else:
        raise ValueError("Only .pptx and .pdf are supported")


if __name__ == "__main__":
    import sys
    stamp(sys.argv[1], sys.argv[2])
